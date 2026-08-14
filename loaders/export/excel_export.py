from datetime import date, datetime
from io import BytesIO

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from loaders.db import get_engine


# Columns to keep and their display names
COLUMNS_STAGES = {
    "external_id":          "ID",
    "name":                 "Nom",
    "partner_name":         "Client / Prospect",
    "user_id":              "Collaborateur affecté",
    "stage_id":             "Etat",
    "x_myco_contract_type": "Type contrat",
    "x_myco_sector_id":     "Secteur",
    "expected_revenue":     "Montant potentiel",
    "date_deadline":        "Date deadline",
}

COLUMNS_WON = {
    **COLUMNS_STAGES,
    "date_closed":          "Date clôture",
}

COLUMNS_LOST = {
    **COLUMNS_STAGES,
    "date_closed":          "Date clôture",
    "lost_reason_id":       "Raison perte",
}

SHEET_CONFIG = [
    {"sheet": "Qualification",           "stage": "Qualification"},
    {"sheet": "RDV 1",                   "stage": "RDV 1"},
    {"sheet": "Preparation proposition", "stage": "Préparation proposition"},
    {"sheet": "Proposition envoyée",     "stage": "Proposition envoyée"},
    {"sheet": "En cours de négociation",    "stage": "En cours de négociation"},
    {"sheet": "Accord de principe",      "stage": "Accord de principe"},
    {"sheet": "Gagne",                   "stage": "Gagné",     "won_lost": True, "cols": "won"},
    {"sheet": "Perdu",                   "stage": "Perdu",     "won_lost": True, "cols": "lost"},
    {"sheet": "Suspendue",                "stage": "Suspendue"},
]

# Stage order and abbreviated labels for the "Recap prorata" sheet columns
RECAP_STAGES = [
    ("Qualification",           "Qualif"),
    ("RDV 1",                   "RDV1"),
    ("Préparation proposition", "Prep prop"),
    ("Proposition envoyée",     "Prop envoyée"),
    ("En cours de négociation", "En cours négo"),
    ("Accord de principe",      "Accord de principe"),
    ("Gagné",                   "Gagné"),
    ("Suspendue",               "Suspendu"),
    ("Perdu",                   "Perdu"),
]

# Stages included in the recap's "Total" column — active pipeline only, Suspendu/Perdu excluded
RECAP_TOTAL_STAGES = {stage for stage, _ in RECAP_STAGES} - {"Suspendue", "Perdu"}

# Theoretical conversion rate applied to each pipeline stage
STAGE_PRORATA = {
    "Qualification":           0.03,
    "RDV 1":                   0.13,
    "Préparation proposition": 0.18,
    "Proposition envoyée":     0.42,
    "En cours de négociation": 0.60,
    "Accord de principe":      0.80,
    "Gagné":                   1.00,
    "Perdu":                   0.0,
    "Suspendue":               0.0,
}

# Pipeline order used by _compute_stage_survival, mirrors the Proba_Survie_Cumul DAX measure
STAGE_ORDER = {
    "Qualification":           1,
    "RDV 1":                   2,
    "Préparation proposition": 3,
    "Proposition envoyée":     4,
    "En cours de négociation": 5,
    "Accord de principe":      6,
    "Gagné":                   7,
    "Perdu":                   8,
}


def _load_opportunities() -> pd.DataFrame:
    """Load all opportunities from raw.opportunities (one row per external_id)."""
    query = "SELECT * FROM raw.opportunities"
    engine = get_engine()
    with engine.connect() as conn:
        return pd.read_sql(query, conn)


def _filter_daily(df: pd.DataFrame) -> pd.DataFrame:
    """
    Filter for donnees_source :
    - date_deadline >= Jan 1 of current year
    - OR (date_closed in current year AND won/lost)
    """
    year = date.today().year
    df["date_deadline"] = pd.to_datetime(df["date_deadline"], errors="coerce")
    df["date_closed"] = pd.to_datetime(df["date_closed"], errors="coerce")

    mask = (
        (df["date_deadline"] >= f"{year}-01-01")
        | (
            df["stage_id"].isin(["Gagné", "Perdu"])
            & (df["date_closed"].dt.year == year)
        )
    )
    return df[mask].copy()


def _filter_monthly(df: pd.DataFrame) -> pd.DataFrame:
    """
    Filter for monthly report (end of month):
    - Open pipeline stages (all but Gagne/Perdu): date_deadline >= Jan 1 of current year
    - Gagne/Perdu: date_closed in current month
    """
    today = date.today()
    year, month = today.year, today.month
    df["date_deadline"] = pd.to_datetime(df["date_deadline"], errors="coerce")
    df["date_closed"] = pd.to_datetime(df["date_closed"], errors="coerce")

    is_won_lost = df["stage_id"].isin(["Gagné", "Perdu"])
    stages_mask = (
        (df["date_deadline"] >= f"{year}-01-01")
        & (df["date_deadline"] <= f"{year}-12-31")
        & ~is_won_lost
    )
    won_lost_mask = (
        is_won_lost
        & (df["date_closed"].dt.year == year)
        & (df["date_closed"].dt.month == month)
    )
    return df[stages_mask | won_lost_mask].copy()


def _stage_sheet_df(df: pd.DataFrame, config: dict, year: int) -> pd.DataFrame:
    """Filter and format the rows/columns for a single stage sheet."""
    if config.get("won_lost"):
        sheet_df = df[
            (df["stage_id"] == config["stage"])
            & (df["date_closed"].dt.year == year)
        ].copy()
        cols = COLUMNS_LOST if config["cols"] == "lost" else COLUMNS_WON
    else:
        sheet_df = df[
            (df["stage_id"] == config["stage"])
            & (df["date_deadline"] >= f"{year}-01-01")
        ].copy()
        cols = COLUMNS_STAGES

    available = {k: v for k, v in cols.items() if k in sheet_df.columns}
    return sheet_df[list(available.keys())].rename(columns=available)


def _as_table(ws, n_rows: int, n_cols: int, sheet_name: str) -> None:
    """Format a worksheet range as an Excel Table with row stripes."""
    if n_cols == 0:
        return
    ref = f"A1:{get_column_letter(n_cols)}{max(n_rows, 1) + 1}"  # +1 for header row
    table_name = "T_" + "".join(c if c.isalnum() else "_" for c in sheet_name)
    tbl = Table(displayName=table_name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(tbl)


def _compute_kpis(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Compute conversion rate and upsell rate from the full (unfiltered) dataset.

    Returns:
        df_kpis: summary table (one row per metric)
        df_upsell_detail: list of clients counted in the upsell numerator
    """
    year = date.today().year

    df = df.copy()
    df["create_date"] = pd.to_datetime(df["create_date"], errors="coerce")
    df["date_closed"] = pd.to_datetime(df["date_closed"], errors="coerce")

    # Taux de conversion = deals gagnés en N / opps créées en N
    won_n = df[(df["stage_id"] == "Gagné") & (df["date_closed"].dt.year == year)]
    created_n = df[df["create_date"].dt.year == year]
    n_won = len(won_n)
    n_created = len(created_n)
    conversion_rate = n_won / n_created if n_created > 0 else 0

    # Taux upsell = clients gagnés en N-1 qui resignent en N dans le même secteur
    won_n1 = df[
        (df["stage_id"] == "Gagné") & (df["date_closed"].dt.year == year - 1)
    ][["partner_id", "partner_name", "x_myco_sector_id"]].dropna(subset=["partner_id"]).drop_duplicates(subset=["partner_id", "x_myco_sector_id"])

    won_n_clients = won_n[["partner_id", "x_myco_sector_id"]].dropna(subset=["partner_id"]).drop_duplicates()

    upsell = won_n1.merge(won_n_clients, on=["partner_id", "x_myco_sector_id"], how="inner")

    n_clients_n1 = won_n1["partner_id"].nunique()
    n_upsell = upsell["partner_id"].nunique()
    upsell_rate = n_upsell / n_clients_n1 if n_clients_n1 > 0 else 0

    df_kpis = pd.DataFrame([
        {
            "Métrique": "Taux de conversion",
            "Valeur": f"{conversion_rate:.1%}",
            "Numérateur": f"{n_won} deals gagnés en {year}",
            "Dénominateur": f"{n_created} opportunités créées en {year}",
            "Méthode de calcul": (
                f"Nb de deals passés en 'Gagné' (date_closed en {year}) "
                f"/ nb total d'opportunités créées cette année (create_date en {year})."
            ),
        },
        {
            "Métrique": "Taux upsell",
            "Valeur": f"{upsell_rate:.1%}",
            "Numérateur": f"{n_upsell} clients ayant resigné dans le même secteur en {year}",
            "Dénominateur": f"{n_clients_n1} clients distincts gagnés en {year - 1}",
            "Méthode de calcul": (
                f"Clients (partner_id) ayant un deal 'Gagné' en {year - 1} "
                f"ET un deal 'Gagné' en {year} dans le même secteur (x_myco_sector_id). "
                f"Dénominateur = nb de clients distincts gagnés en {year - 1}."
            ),
        },
    ])

    df_upsell_detail = (
        upsell[["partner_id", "partner_name", "x_myco_sector_id"]]
        .rename(columns={
            "partner_id":       "ID Client",
            "partner_name":     "Client",
            "x_myco_sector_id": "Secteur",
        })
        .sort_values("Client")
        .reset_index(drop=True)
    )

    return df_kpis, df_upsell_detail


# Sheet names used before the 2026-07-22 rename (d1ad61f) — kept so écart vs S-2 still
# resolves against files generated before that fix. Safe to drop once no S-2 candidate
# predates it (two weekly runs after the rename).
LEGACY_SHEET_NAMES = {
    "Proposition envoyée":     "Proposition envoyee",
    "En cours de négociation": "En cours negociation",
}


def _stage_totals_from_workbook(content: BytesIO) -> dict[str, dict[str, float]]:
    """Sum 'Montant potentiel' and count rows per stage sheet from a previously generated workbook."""
    totals = {}
    xls = pd.ExcelFile(content)
    for config in SHEET_CONFIG:
        sheet_name = config["sheet"]
        if sheet_name not in xls.sheet_names:
            sheet_name = LEGACY_SHEET_NAMES.get(config["stage"])
            if sheet_name not in xls.sheet_names:
                continue
        sheet_df = xls.parse(sheet_name)
        totals[config["stage"]] = {
            "montant": sheet_df["Montant potentiel"].sum() if "Montant potentiel" in sheet_df.columns else None,
            "nb": len(sheet_df),
        }
    return totals


def _previous_stage_totals() -> dict[str, dict[str, float]]:
    """Fetch montant/count per stage from the donnees_source uploaded to SharePoint two weeks ago (S-2).

    Returns an empty dict (no écart computed) if there's no S-2 file yet or the
    SharePoint fetch fails for any reason (e.g. missing credentials during local runs).
    """
    try:
        from .sharepoint_upload import download_previous_donnees_source
        content = download_previous_donnees_source()
    except Exception as e:
        print(f"Could not fetch previous donnees_source for écart comparison: {e}")
        return {}
    if content is None:
        return {}
    return _stage_totals_from_workbook(content)


def _compute_stage_survival(df_all: pd.DataFrame) -> dict[str, float]:
    """Cumulative survival probability per stage, mirroring the Proba_Survie_Cumul DAX measure.

    For each pipeline stage (Qualification..Accord de principe), the loss probability is the
    share of opportunities that ever reached at least that stage (current stage, across the full
    unfiltered dataset) and ended up "Perdu" for a real reason (blank or "Fermé - Plus une opp" —
    other lost reasons are excluded from the loss count, same as the DAX measure). The survival
    probability for a stage is the product of (1 - loss probability) over that stage and every
    stage ahead of it up to "Accord de principe".
    """
    ordre = df_all["stage_id"].map(STAGE_ORDER)
    is_real_loss = (df_all["stage_id"] == "Perdu") & (
        df_all["lost_reason_id"].isna() | (df_all["lost_reason_id"] == "Fermé - Plus une opp")
    )

    proba_perte = {}
    for stage, s in STAGE_ORDER.items():
        if s >= 7:  # Gagné / Perdu are terminal outcomes, not pipeline stages
            continue
        reached = ordre >= s
        opps_a_etape = int(reached.sum())
        opps_perdues = int((reached & is_real_loss).sum())
        proba_perte[s] = opps_perdues / opps_a_etape if opps_a_etape else 0.0

    survival = {"Gagné": 1.0, "Perdu": 0.0}
    for stage, s in STAGE_ORDER.items():
        if stage in survival:
            continue
        prod = 1.0
        for o in range(s, 7):
            prod *= 1 - proba_perte[o]
        survival[stage] = prod
    return survival


def _compute_recap(df: pd.DataFrame, df_all: pd.DataFrame) -> pd.DataFrame:
    """Build the recap: montants, écart vs export précédent, prorata réel/théorique — in K€.

    One column per stage (abbreviated labels, Suspendu/Perdu last) plus a "Total" column
    limited to the active pipeline (Qualification..Gagné), one row per indicator.
    """
    year = date.today().year
    previous_totals = _previous_stage_totals()
    survival = _compute_stage_survival(df_all)
    config_by_stage = {c["stage"]: c for c in SHEET_CONFIG}

    nb_contrats, ecarts_nb, montants, ecarts, proratas_reels, proratas_theo, montants_theo = (
        {}, {}, {}, {}, {}, {}, {}
    )
    for stage, label in RECAP_STAGES:
        sheet_df = _stage_sheet_df(df, config_by_stage[stage], year)
        montant_total = (
            sheet_df["Montant potentiel"].sum() / 1000
            if "Montant potentiel" in sheet_df.columns else 0.0
        )
        nb_total = len(sheet_df)
        precedent = previous_totals.get(stage)
        montant_precedent = precedent["montant"] if precedent else None
        nb_precedent = precedent["nb"] if precedent else None
        ecart = montant_total - montant_precedent / 1000 if montant_precedent is not None else None
        ecart_nb = nb_total - nb_precedent if nb_precedent is not None else None
        prorata_theo = STAGE_PRORATA.get(stage, 0.0)
        prorata_reel = survival.get(stage, 0.0)

        nb_contrats[label] = nb_total
        ecarts_nb[label] = ecart_nb
        montants[label] = round(montant_total)
        ecarts[label] = round(ecart) if ecart is not None else None
        proratas_reels[label] = f"{prorata_reel:.1%}"
        proratas_theo[label] = f"{prorata_theo:.0%}"
        montants_theo[label] = round(montant_total * prorata_theo, 2)

    total_nb_contrats = sum(nb_contrats[label] for stage, label in RECAP_STAGES if stage in RECAP_TOTAL_STAGES)
    total_montants = round(sum(montants[label] for stage, label in RECAP_STAGES if stage in RECAP_TOTAL_STAGES))
    total_montants_theo = round(
        sum(montants_theo[label] for stage, label in RECAP_STAGES if stage in RECAP_TOTAL_STAGES), 2
    )

    rows = [
        {"Indicateur": "Nombre de contrats",     **nb_contrats,    "Total": total_nb_contrats},
        {"Indicateur": "Ecart nombre vs S-2",    **ecarts_nb,      "Total": None},
        {"Indicateur": "Montants",              **montants,       "Total": f"{total_montants}K€"},
        {"Indicateur": "Ecart vs S-2",           **ecarts,         "Total": None},
        {"Indicateur": "Prorata réel",           **proratas_reels, "Total": None},
        {"Indicateur": "Prorata théo",           **proratas_theo,  "Total": None},
        {"Indicateur": "Montants prorata théo",  **montants_theo,  "Total": f"{total_montants_theo}K€"},
    ]
    return pd.DataFrame(rows)


def _build_excel(
    df: pd.DataFrame,
    df_kpis: pd.DataFrame | None = None,
    df_upsell_detail: pd.DataFrame | None = None,
    df_recap: pd.DataFrame | None = None,
) -> BytesIO:
    """Build an Excel file in memory with one sheet per stage (+ optional summary sheets), each formatted as a Table."""
    year = date.today().year
    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # Summary sheets first
        if df_recap is not None:
            df_recap.to_excel(writer, sheet_name="Recap prorata", index=False)
            ws = writer.sheets["Recap prorata"]
            _as_table(ws, len(df_recap), len(df_recap.columns), "Recap prorata")

        if df_kpis is not None:
            df_kpis.to_excel(writer, sheet_name="KPIs", index=False)
            ws = writer.sheets["KPIs"]
            _as_table(ws, len(df_kpis), len(df_kpis.columns), "KPIs")

        if df_upsell_detail is not None:
            df_upsell_detail.to_excel(writer, sheet_name="Upsell - Detail", index=False)
            ws = writer.sheets["Upsell - Detail"]
            _as_table(ws, len(df_upsell_detail), len(df_upsell_detail.columns), "Upsell - Detail")

        for config in SHEET_CONFIG:
            sheet_df = _stage_sheet_df(df, config, year)
            sheet_df.to_excel(writer, sheet_name=config["sheet"], index=False)

            ws = writer.sheets[config["sheet"]]
            _as_table(ws, len(sheet_df), len(sheet_df.columns), config["sheet"])

    output.seek(0)
    return output


def generate_donnees_source() -> tuple[BytesIO, str]:
    """Generate donnees_source_YYYY_MM_DD.xlsx."""
    df_all = _load_opportunities()
    df = _filter_daily(df_all)
    df_recap = _compute_recap(df, df_all)
    content = _build_excel(df, df_recap=df_recap)
    filename = f"donnees_source_{datetime.now().strftime('%Y_%m_%d_%H%M%S')}.xlsx"
    print(f"Generated {filename} — {len(df)} rows across 10 sheets")
    return content, filename


def generate_reporting_mensuel() -> tuple[BytesIO, str]:
    """Generate RM-MM-YY.xlsx."""
    df_all = _load_opportunities()
    df_kpis, df_upsell_detail = _compute_kpis(df_all)
    df = _filter_monthly(df_all)
    content = _build_excel(df, df_kpis, df_upsell_detail)
    filename = f"RM-{date.today().strftime('%b-%y')}.xlsx"
    print(f"Generated {filename} — {len(df)} rows across 11 sheets")
    return content, filename
